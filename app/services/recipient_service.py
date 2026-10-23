import csv
import io
import os
import re
from typing import List, Optional, Tuple, Set
from email_validator import validate_email, EmailNotValidError

from app import config

# Loose pattern used to find candidate addresses inside free text (PDF / Word / TXT).
# Every match is still checked with email_validator before being accepted.
EMAIL_PATTERN = re.compile(r"[A-Za-z0-9._%+'-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}")

SUPPORTED_EXTENSIONS = {".csv", ".txt", ".pdf", ".docx"}


def _add_candidate(raw: str, label: str, valid: List[str], invalid: List[str], seen: Set[str]) -> None:
    email = raw.strip().strip("<>;,\"'")
    if not email:
        return
    try:
        email = validate_email(email, check_deliverability=False).normalized
    except EmailNotValidError as e:
        invalid.append(f"{label}: {raw.strip()} ({e})")
        return
    if email.lower() not in seen:
        seen.add(email.lower())
        valid.append(email)


def validate_recipients(csv_content: str) -> Tuple[List[str], List[str]]:
    """
    Parses CSV content and validates email addresses.
    If an 'email' column exists only that column is used; otherwise every
    cell containing '@' (including the first row) is treated as a recipient.
    Returns a tuple of (valid_emails, invalid_rows).
    """
    valid_emails: List[str] = []
    invalid_rows: List[str] = []
    seen_emails: Set[str] = set()

    rows = list(csv.reader(io.StringIO(csv_content.strip())))
    if not rows:
        return [], ["Empty file"]

    header = [col.strip().lower() for col in rows[0]]
    if "email" in header:
        email_idx = header.index("email")
        for row_idx, row in enumerate(rows[1:], start=2):
            if len(row) > email_idx:
                _add_candidate(row[email_idx], f"Row {row_idx}", valid_emails, invalid_rows, seen_emails)
    else:
        for row_idx, row in enumerate(rows, start=1):
            for cell in row:
                if "@" in cell:
                    _add_candidate(cell, f"Row {row_idx}", valid_emails, invalid_rows, seen_emails)

    return valid_emails, invalid_rows


def extract_emails_from_text(text: str) -> List[str]:
    """Finds, validates and de-duplicates all email addresses in free text."""
    valid: List[str] = []
    invalid: List[str] = []
    seen: Set[str] = set()
    for match in EMAIL_PATTERN.findall(text):
        _add_candidate(match.rstrip("."), "Text", valid, invalid, seen)
    return valid


def parse_email_list(text: str) -> Tuple[List[str], List[str]]:
    """Parses addresses separated by commas, semicolons or whitespace (e.g. the CC box)."""
    valid: List[str] = []
    invalid: List[str] = []
    seen: Set[str] = set()
    for item in re.split(r"[,;\s]+", text or ""):
        if item:
            _add_candidate(item, "Address", valid, invalid, seen)
    return valid, invalid


def _pdf_text(content: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(content: bytes) -> str:
    import docx
    document = docx.Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore")


def _file_text(filename: str, ext: str, content: bytes) -> str:
    try:
        if ext == ".pdf":
            return _pdf_text(content)
        if ext == ".docx":
            return _docx_text(content)
        return _decode(content)
    except Exception as e:
        raise ValueError(f"Could not read {filename}: {e}") from e


def extract_recipients(filename: str, content: bytes, channel: str = "email") -> Tuple[List[str], List[str]]:
    """
    Extracts recipients from an uploaded file (CSV, TXT, PDF or DOCX).
    channel "email" returns email addresses, "whatsapp" returns phone numbers.
    Returns a tuple of (recipients, problems).
    """
    ext = os.path.splitext(filename or "")[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type '{ext or filename}'. Use CSV, TXT, PDF or DOCX.")

    if channel == "whatsapp":
        if ext == ".csv":
            return phones_from_csv(_decode(content))
        phones = extract_phones_from_text(_file_text(filename, ext, content))
        return phones, [] if phones else [f"No phone numbers found in {filename}"]

    if ext == ".csv":
        return validate_recipients(_decode(content))
    emails = extract_emails_from_text(_file_text(filename, ext, content))
    problems = [] if emails else [f"No email addresses found in {filename}"]
    return emails, problems


# ---------------------------------------------------------------- phone numbers

PHONE_PATTERN = re.compile(r"(?<![\w+])(?:\+|00)?\d[\d \-().]{7,17}\d(?!\w)")
PHONE_HEADERS = ("phone", "mobile", "whatsapp", "contact", "number", "cell", "tel")


def normalize_phone(raw: str, strict: bool = False) -> Optional[str]:
    """
    Converts a phone number to international format (+<country><number>).
    10-digit numbers get the default country code. With strict=True (used for
    free text from PDFs etc.) only number shapes that are clearly phone numbers
    are accepted, so invoice numbers, dates and PIN codes are ignored.
    """
    raw = raw.strip()
    digits = re.sub(r"\D", "", raw)
    international = raw.startswith("+") or raw.startswith("00")
    if raw.startswith("00"):
        digits = digits[2:]
    cc = config.WA_DEFAULT_COUNTRY_CODE

    if international:
        return "+" + digits if 8 <= len(digits) <= 15 else None
    if len(digits) == 11 and digits.startswith("0"):
        digits = digits[1:]
    if len(digits) == 10:
        if cc == "91" and digits[0] not in "6789":
            return None
        return "+" + cc + digits
    if len(digits) == len(cc) + 10 and digits.startswith(cc):
        return "+" + digits
    if not strict and 11 <= len(digits) <= 15 and not digits.startswith("0"):
        return "+" + digits
    return None


def _add_phone(raw: str, label: str, strict: bool, valid: List[str], invalid: List[str], seen: Set[str]) -> None:
    phone = normalize_phone(raw, strict)
    if phone is None:
        if not strict:
            invalid.append(f"{label}: {raw.strip()} (not a valid phone number)")
        return
    if phone not in seen:
        seen.add(phone)
        valid.append(phone)


def extract_phones_from_text(text: str) -> List[str]:
    valid: List[str] = []
    seen: Set[str] = set()
    for match in PHONE_PATTERN.findall(text):
        _add_phone(match, "Text", True, valid, [], seen)
    return valid


def phones_from_csv(csv_content: str) -> Tuple[List[str], List[str]]:
    """Uses phone/mobile/whatsapp columns if present, otherwise scans every cell."""
    valid: List[str] = []
    invalid: List[str] = []
    seen: Set[str] = set()
    rows = list(csv.reader(io.StringIO(csv_content.strip())))
    if not rows:
        return [], ["Empty file"]

    header = [col.strip().lower() for col in rows[0]]
    phone_cols = [i for i, col in enumerate(header) if any(h in col for h in PHONE_HEADERS)]
    if phone_cols:
        for row_idx, row in enumerate(rows[1:], start=2):
            for i in phone_cols:
                if i < len(row) and row[i].strip():
                    _add_phone(row[i], f"Row {row_idx}", False, valid, invalid, seen)
    else:
        for row in rows:
            for cell in row:
                for match in PHONE_PATTERN.findall(cell):
                    _add_phone(match, "Text", True, valid, invalid, seen)
    if not valid and not invalid:
        invalid.append("No phone numbers found")
    return valid, invalid


def parse_phone_list(text: str) -> Tuple[List[str], List[str]]:
    """Parses numbers separated by new lines, commas or semicolons."""
    valid: List[str] = []
    invalid: List[str] = []
    seen: Set[str] = set()
    for item in re.split(r"[,;\n]+", text or ""):
        if item.strip():
            _add_phone(item, "Number", False, valid, invalid, seen)
    return valid, invalid
